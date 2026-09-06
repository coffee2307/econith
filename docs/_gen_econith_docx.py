# -*- coding: utf-8 -*-
"""Generate ECONITH_Project_Documentation.docx — KHKT format (A4, TNR 14, 15 pages)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Palette (giữ phong cách whitepaper)
BLACK = RGBColor(0x11, 0x11, 0x11)
DARK = RGBColor(0x2A, 0x2A, 0x2E)
GRAY = RGBColor(0x5A, 0x5A, 0x62)
MUTED = RGBColor(0x7A, 0x7A, 0x82)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
LINE = "D0D5DD"
SOFT_BG = "F5F7FA"

FONT = "Times New Roman"
BODY = 14
OUT = Path(__file__).resolve().parent.parent / "ECONITH_Project_Documentation.docx"


def set_run_font(run, *, name=FONT, size=BODY, bold=False, italic=False, color=BLACK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_para(
    doc,
    text="",
    *,
    style=None,
    size=BODY,
    bold=False,
    italic=False,
    color=BLACK,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    space_before=0,
    space_after=4,
    first_line_indent=None,
):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.0
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_runs(p, parts):
    for text, kwargs in parts:
        run = p.add_run(text)
        set_run_font(run, **kwargs)
    return p


def h1(doc, text):
    p = add_para(doc, text, size=BODY, bold=True, color=ACCENT, space_before=10, space_after=4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def h2(doc, text):
    return add_para(doc, text, size=BODY, bold=True, color=DARK, space_before=8, space_after=3)


def h3(doc, text):
    return add_para(doc, text, size=BODY, bold=True, italic=True, color=DARK, space_before=6, space_after=2)


def body(doc, text, *, indent=True):
    return add_para(doc, text, size=BODY, space_after=4, first_line_indent=0.75 if indent else None)


def bullet(doc, text, *, level=0):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.left_indent = Cm(0.75 + level * 0.4)
    run = p.add_run(text)
    set_run_font(run, size=BODY, color=BLACK)
    return p


def quote(doc, text):
    p = add_para(doc, text, italic=True, size=BODY, color=DARK, space_before=4, space_after=6)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(0.3)
    return p


def set_cell_shading(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color=LINE, sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def fill_cell(cell, text, *, bold=False, size=12, color=BLACK, align=WD_ALIGN_PARAGRAPH.LEFT, bg=None):
    """Table cell: 12pt to fit table density within TNR report (body still 14)."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(1)
    pf.space_after = Pt(1)
    pf.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    if bg:
        set_cell_shading(cell, bg)
    set_cell_borders(cell)


def make_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        fill_cell(
            table.rows[0].cells[i],
            h,
            bold=True,
            size=12,
            color=RGBColor(0xFF, 0xFF, 0xFF),
            bg="1F4E79",
        )
    for r_i, row in enumerate(rows):
        bg = SOFT_BG if r_i % 2 else None
        for c_i, val in enumerate(row):
            fill_cell(table.rows[r_i + 1].cells[c_i], val, size=12, bg=bg)
    add_para(doc, "", space_after=4)
    return table


def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), LINE)
    pBdr.append(bottom)
    pPr.append(pBdr)


def setup_page(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(2)
    run = hp.add_run("ECONITH  ·  Báo cáo dự án nghiên cứu")
    set_run_font(run, size=11, color=MUTED)
    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Trang ")
    set_run_font(run, size=11, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run2 = fp.add_run()
    run2._r.append(fld_begin)
    run2._r.append(instr)
    run2._r.append(fld_end)
    set_run_font(run2, size=11, color=MUTED)
    run3 = fp.add_run(" / 15")
    set_run_font(run3, size=11, color=MUTED)


def page_break(doc):
    doc.add_page_break()


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------


def cover_page(doc):
    add_para(doc, "", space_after=8)
    add_para(
        doc,
        "ECONITH",
        size=26,
        bold=True,
        color=ACCENT,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=2,
    )
    add_para(
        doc,
        "Economic Intelligence Through Simulation & Quant Research",
        size=12,
        italic=True,
        color=GRAY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=8,
    )
    add_hr(doc)
    add_para(
        doc,
        "NỀN TẢNG NGHIÊN CỨU MÔ PHỎNG TÀI CHÍNH — KINH TẾ",
        size=14,
        bold=True,
        color=BLACK,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
    )
    add_para(
        doc,
        "Agent-Based Macro Simulation kết hợp AI Quant\n"
        "để nghiên cứu lan truyền shock vĩ mô tới thị trường tài chính",
        size=14,
        color=DARK,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=10,
    )
    add_para(
        doc,
        "BÁO CÁO DỰ ÁN NGHIÊN CỨU / KỸ THUẬT",
        size=14,
        bold=True,
        color=GRAY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )

    meta = [
        ("Phiên bản tài liệu", "1.1"),
        ("Ngày phát hành", date.today().strftime("%d/%m/%Y")),
        ("Định dạng", "A4 · Times New Roman 14 · giãn dòng 1.0 · ≤ 15 trang"),
        ("Phạm vi", "Kiến trúc · Phương pháp · Trạng thái kỹ thuật theo codebase"),
        ("Chế độ vận hành chính", "SIM / DEMO (nghiên cứu & thử nghiệm)"),
        ("Phân loại", "Research Platform — không phải hệ giao dịch tự động"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(meta):
        fill_cell(table.rows[i].cells[0], k, bold=True, size=12, color=DARK, bg=SOFT_BG)
        fill_cell(table.rows[i].cells[1], v, size=12, color=BLACK)

    add_para(doc, "", space_after=10)
    quote(
        doc,
        "ECONITH là nền tảng nghiên cứu mô phỏng tài chính–kinh tế sử dụng mô hình đa tác nhân "
        "(agent-based simulation) kết hợp AI Quant để nghiên cứu sự lan truyền của các biến động "
        "kinh tế vĩ mô tới thị trường tài chính — trong môi trường có kiểm soát, có thể tái lập, "
        "và tách biệt rõ ràng giữa mô phỏng và thực thi thực tế.",
    )
    add_para(
        doc,
        "Tài liệu mô tả bản chất hệ thống dựa trên codebase hiện tại; không tuyên bố vượt quá "
        "khả năng đã triển khai. Không ghi thông tin đơn vị theo quy định trình bày báo cáo.",
        size=12,
        italic=True,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=8,
    )


def toc_page(doc):
    h1(doc, "MỤC LỤC")
    items = [
        ("01", "Tóm tắt dự án"),
        ("02", "Tổng quan về đề tài"),
        ("03", "Vấn đề nghiên cứu"),
        ("04", "Mục tiêu dự án"),
        ("05", "Kiến trúc hệ thống"),
        ("06", "ECONITH World"),
        ("07", "ECONITH Quant"),
        ("08", "Cơ chế kết nối World ↔ Quant"),
        ("09", "Phương pháp nghiên cứu & thực nghiệm"),
        ("10", "Dữ liệu & thành phần kỹ thuật"),
        ("11", "Tính sáng tạo"),
        ("12", "Ứng dụng thực tiễn"),
        ("13", "Trạng thái hiện tại & giới hạn"),
        ("14", "Hướng phát triển"),
        ("15", "Kết luận"),
        ("A", "Phụ lục — Bản đồ module & thuật ngữ"),
    ]
    for num, title in items:
        p = add_para(doc, "", space_after=2)
        add_runs(
            p,
            [
                (f"{num}    ", {"size": BODY, "bold": True, "color": ACCENT}),
                (title, {"size": BODY, "color": BLACK}),
            ],
        )


def section_summary(doc):
    h1(doc, "01  ·  TÓM TẮT DỰ ÁN")

    h2(doc, "1.1. Định vị")
    body(
        doc,
        "ECONITH là nền tảng nghiên cứu (research platform) gồm hai miền chính: "
        "ECONITH World — mô phỏng kinh tế vĩ mô đa tác nhân; và ECONITH Quant — pipeline AI định lượng "
        "để quan sát phản ứng thị trường dưới các điều kiện khác nhau. Hệ thống chạy trên kiến trúc "
        "event-driven với chế độ REALITY / SIMULATION được cô lập để tránh nhiễm dữ liệu mô phỏng vào "
        "đường thực thi thực tế.",
    )

    h2(doc, "1.2. Tính mới")
    bullet(doc, "Tích hợp mô phỏng agent-based macro với pipeline AI Quant trong một runtime thống nhất.")
    bullet(doc, "Cơ chế coupling hai chiều macro ↔ micro (cross-impact), bật/tắt theo chế độ nghiên cứu.")
    bullet(doc, "Vòng Hypothesis Runner: sinh giả thuyết shock → chạy kịch bản → đo delta → sealed rollout.")
    bullet(doc, "Governance rủi ro độc lập (Sentinel) và minh bạch trạng thái thực thi (LIVE / SYNTHETIC / DEGRADED).")
    bullet(doc, "Không claim “dự đoán tương lai” hay “digital twin hoàn chỉnh”; tập trung môi trường thí nghiệm có kiểm soát.")

    h2(doc, "1.3. Tính khoa học")
    bullet(doc, "Phương pháp R&D: thu thập dữ liệu → xây hệ thống → calibration → so sánh baseline → ablation.")
    bullet(doc, "Hướng đánh giá: historical replay, scenario simulation, moment-matching, reproducibility.")
    bullet(doc, "Biến độc lập/phụ thuộc đo được: coupling ON/OFF, loại shock, MAE biến động, risk metrics, moment error.")
    bullet(doc, "Kết quả âm (coupling không cải thiện baseline) vẫn được xem là kết quả nghiên cứu hợp lệ.")

    h2(doc, "1.4. Tính thực tiễn")
    bullet(doc, "Hỗ trợ nghiên cứu viên / người học phân tích kịch bản shock vĩ mô và quan sát phản ứng thị trường.")
    bullet(doc, "Phục vụ stress testing có chủ đích — khác backtest chỉ replay lịch sử giá.")
    bullet(doc, "Dashboard operator quan sát World state, Quant cockpit và nhật ký sự kiện hệ thống.")
    bullet(doc, "Chế độ SIM/DEMO phù hợp giáo dục và nghiên cứu; không yêu cầu vốn thật để chạy thí nghiệm.")

    h2(doc, "1.5. Phạm vi không bao gồm")
    bullet(doc, "Không phải trading bot kiếm lời tự động.")
    bullet(doc, "Không phải hệ thống thay thế chuyên gia tài chính hoặc nhà hoạch định chính sách.")
    bullet(doc, "Không phải digital twin đã hiệu chỉnh đầy đủ cho toàn bộ nền kinh tế thế giới.")
    bullet(doc, "Không tuyên bố AI dự đoán chính xác tương lai thị trường.")


def section_overview(doc):
    h1(doc, "02  ·  TỔNG QUAN VỀ ĐỀ TÀI")

    h2(doc, "2.1. ECONITH là gì?")
    body(
        doc,
        "ECONITH (Economic Intelligence Through Simulation & Quant Research) là hệ thống phần mềm "
        "nghiên cứu cho phép đặt giả thuyết về cú sốc kinh tế vĩ mô (lãi suất, lạm phát, thuế quan…), "
        "chạy mô phỏng lan truyền qua mạng lưới quốc gia–tác nhân, và quan sát cách tín hiệu thị trường "
        "vi mô phản ứng trong cùng một khung thời gian thí nghiệm.",
    )
    body(
        doc,
        "Người dùng chính không phải “nhà giao dịch tự động”, mà là nghiên cứu viên, kỹ sư dữ liệu "
        "và người học muốn kiểm tra ý tưởng trong môi trường có thể tái lập. Phát biểu một câu: "
        "ECONITH nghiên cứu vấn đề thiếu môi trường thí nghiệm đo được cho lan truyền shock vĩ mô "
        "tới thị trường, bằng phương pháp agent-based simulation kết hợp AI Quant và so sánh "
        "baseline/ablation, nhằm hỗ trợ phân tích rủi ro và kiểm chứng giả thuyết trong điều kiện tái lập được.",
    )

    h2(doc, "2.2. Vì sao dự án được tạo ra?")
    body(doc, "Trong thực tiễn phân tích tài chính–kinh tế, các công cụ thường tách rời:")
    bullet(doc, "Phân tích vĩ mô / kịch bản chính sách (bảng tính, mô hình đơn giản, narrative).")
    bullet(doc, "Mô phỏng agent-based kinh tế (thường không nối microstructure thị trường).")
    bullet(doc, "Backtesting và AI Quant trên dữ liệu giá (thường thiếu shock vĩ mô có cấu trúc).")
    bullet(doc, "Hệ thống rủi ro / thực thi (thường tách khỏi vòng nghiên cứu giả thuyết).")
    body(
        doc,
        "Hệ quả: khó trả lời câu hỏi dạng “Nếu xảy ra shock X, tín hiệu thị trường và metric rủi ro "
        "thay đổi thế nào so với baseline và so với shock ngẫu nhiên?” trong một protocol thống nhất.",
    )

    h2(doc, "2.3. Bối cảnh vấn đề")
    body(
        doc,
        "Thị trường tài sản số phản ứng nhanh với tin tức và điều kiện vĩ mô, nhưng phương pháp đánh giá "
        "tín hiệu phổ biến vẫn dựa chủ yếu vào lịch sử giá. Ngược lại, mô phỏng kinh tế thường dừng ở tầng "
        "biến số vĩ mô mà không có cầu nối đo được sang biến động vi cấu trúc (order-flow, volatility "
        "perceived, regime). ECONITH được thiết kế như phòng thí nghiệm để nối hai tầng này — với nhận thức "
        "rõ rằng mô phỏng là stylized và cần calibration / validation, không phải bản sao trung thực của thế giới thực.",
    )


def section_problem(doc):
    h1(doc, "03  ·  VẤN ĐỀ NGHIÊN CỨU")

    h2(doc, "3.1. Phát biểu vấn đề")
    quote(
        doc,
        "Thiếu một môi trường mô phỏng có khả năng nghiên cứu sự lan truyền của các cú sốc kinh tế "
        "vĩ mô tới hành vi thị trường tài chính trong điều kiện có thể kiểm soát, đo lường và tái lập.",
    )

    h2(doc, "3.2. Hạn chế của phương pháp truyền thống")
    make_table(
        doc,
        ["Phương pháp", "Điểm mạnh", "Hạn chế liên quan đến đề tài"],
        [
            ["Backtesting lịch sử", "Dựa dữ liệu thật", "Khó chủ động tạo shock có cấu trúc"],
            ["Mô hình macro đơn lẻ", "Diễn giải chính sách", "Thường không nối microstructure / AI Quant"],
            ["Agent-based kinh tế", "Tương tác đa tác nhân", "Ít pipeline đánh giá phản ứng thị trường số"],
            ["Trading / Quant bot", "Thực thi & tối ưu tín hiệu", "Mục tiêu PnL; thiếu protocol nghiên cứu shock"],
        ],
    )

    h2(doc, "3.3. Khoảng trống ECONITH hướng tới")
    bullet(doc, "Một runtime thống nhất cho mô phỏng World + quan sát Quant.")
    bullet(doc, "Coupling macro→micro đo được (volatility bias, order-flow shock, regime pressure).")
    bullet(doc, "Protocol giả thuyết: generate → simulate → measure → seal.")
    bullet(doc, "Ablation có/không World; so sánh với baseline chỉ market data.")
    bullet(doc, "Isolation mode: mô phỏng không được nhiễm đường thực thi REALITY.")

    h2(doc, "3.4. Câu hỏi nghiên cứu định hướng")
    body(
        doc,
        "RQ1. Việc bổ sung shock macro có cấu trúc từ mô phỏng agent có cải thiện khả năng đánh giá "
        "biến động ngắn hạn so với mô hình chỉ dùng dữ liệu thị trường trên các cửa sổ sự kiện lịch sử không?",
    )
    body(
        doc,
        "RQ2. Calibration moment-matching trên chuỗi macro thực có giảm sai số phân phối mô phỏng "
        "so với hệ số mặc định không?",
    )
    body(
        doc,
        "RQ3. Trong ablation có/không coupling World→Quant, các metric rủi ro (drawdown, số lần "
        "cảnh báo rủi ro) thay đổi thế nào dưới cùng một shock?",
    )

    h2(doc, "3.5. Giả thuyết khoa học")
    bullet(doc, "H1: Trên cửa sổ sự kiện lịch sử, cấu hình có coupling cải thiện đánh giá biến động ngắn hạn so với baseline — và tốt hơn random-shock control.")
    bullet(doc, "H2: Hệ số sau calibration giảm moment error so với hệ số mặc định.")
    bullet(doc, "H3: Ablation coupling ON/OFF tạo khác biệt có hệ thống trên risk metrics dưới cùng shock.")


def section_goals(doc):
    h1(doc, "04  ·  MỤC TIÊU DỰ ÁN")

    h2(doc, "4.1. Mục tiêu tổng quát")
    body(
        doc,
        "Xây dựng và vận hành một nền tảng nghiên cứu cho phép đặt giả thuyết về shock kinh tế vĩ mô, "
        "mô phỏng lan truyền qua hệ đa tác nhân, và đánh giá phản ứng thị trường / rủi ro trong môi trường "
        "thí nghiệm có kiểm soát.",
    )

    h2(doc, "4.2. Mục tiêu kỹ thuật")
    make_table(
        doc,
        ["#", "Mục tiêu", "Thành phần liên quan"],
        [
            ["T1", "Môi trường mô phỏng kinh tế đa tác nhân (hub chọn lọc)", "WorldKernel, Sovereign Graph"],
            ["T2", "Mô hình hóa biến số vĩ mô và quan hệ thương mại / thuế quan", "Macro vectors, Trade matrix"],
            ["T3", "Cơ chế macro → micro và phản hồi micro → macro", "Cross Impact"],
            ["T4", "Pipeline AI Quant: feature, suy luận, đánh giá tín hiệu", "Feature pipeline, Predictor, AIBridge"],
            ["T5", "Vòng thí nghiệm giả thuyết có log tái lập", "Hypothesis Runner, Sealed rollouts"],
            ["T6", "Lớp rủi ro độc lập và cô lập REALITY / SIMULATION", "Sentinel, EventBus mode gates"],
            ["T7", "Thu thập dữ liệu market / macro / tradfi tách runtime", "Collectors, raw lake"],
        ],
    )

    h2(doc, "4.3. Non-goals")
    bullet(doc, "Tối ưu lợi nhuận giao dịch tự động trên vốn thật.")
    bullet(doc, "Mô phỏng chính xác toàn bộ nền kinh tế toàn cầu.")
    bullet(doc, "Thay thế mô hình kinh tế lượng truyền thống (DSGE, VAR) trong hoạch định chính sách.")


def section_architecture(doc):
    h1(doc, "05  ·  KIẾN TRÚC HỆ THỐNG")

    h2(doc, "5.1. Sơ đồ phân lớp")
    quote(
        doc,
        "Data Layer → World Simulation Layer → Macro–Micro Interaction Layer → "
        "Quant Intelligence Layer → Risk & Evaluation Layer → Dashboard",
    )
    make_table(
        doc,
        ["Lớp", "Vai trò", "Module tiêu biểu"],
        [
            ["Data", "Thu thập & raw lake", "collectors/market_coin, macro_global, tradfi_assets"],
            ["World Simulation", "Trạng thái quốc gia, agent, kịch bản", "WorldKernel, SovereignWorldGraph"],
            ["Macro–Micro", "Dịch shock giữa hai miền", "cross_impact (macro_to_micro / quant_to_macro)"],
            ["Quant Intelligence", "Feature, suy luận, tín hiệu", "feature_pipeline, Predictor, AIBridge"],
            ["Risk & Evaluation", "Veto rủi ro, backtest, deploy gate", "Sentinel, evaluation/backtest"],
            ["Dashboard", "Quan sát operator", "Next.js /quant, /world"],
        ],
    )

    h2(doc, "5.2. Runtime lõi")
    bullet(doc, "EventBus — hợp đồng pub/sub giữa các subsystem.")
    bullet(doc, "TimeEngine + TickPipeline (5 pha) — tiến trình mô phỏng/quan sát có nhịp.")
    bullet(doc, "QuantMode REALITY | SIMULATION — ranh giới chủ quyền dữ liệu.")
    bullet(doc, "main.py (FastAPI) — điểm vào backend; dashboard Next.js — bề mặt quan sát.")

    h2(doc, "5.3. Luồng vận hành khái quát")
    body(
        doc,
        "market/macro/tradfi → EventBus → signal → risk gate → execution/fill (DEMO/SYNTHETIC hoặc LIVE khi cấu hình) "
        "→ telemetry → dashboard. Song song: World hypothesis / rollouts → sealed JSONL → (tùy chọn) ingest training.",
    )

    h2(doc, "5.4. Thứ tự lõi theo mục tiêu nghiên cứu")
    bullet(doc, "1) Simulation / World — môi trường thí nghiệm.")
    bullet(doc, "2) Cross-impact — điểm khác biệt phương pháp.")
    bullet(doc, "3) Quant — công cụ đo phản ứng thị trường.")
    bullet(doc, "4) Risk system — ràng buộc an toàn và metric độ bền.")
    bullet(doc, "5) AI training — hỗ trợ; không phải mục tiêu duy nhất của đề tài.")


def section_world(doc):
    h1(doc, "06  ·  ECONITH WORLD")

    h2(doc, "6.1. Bản chất")
    body(
        doc,
        "ECONITH World là lớp mô phỏng kinh tế vĩ mô dạng agent-based, dùng để mô hình hóa các mối quan hệ "
        "kinh tế dưới dạng môi trường nghiên cứu — không nhằm “mô phỏng chính xác nền kinh tế thế giới”.",
    )

    h2(doc, "6.2. Sovereign Graph và agents")
    body(doc, "Mỗi quốc gia/hub trong đồ thị có bốn vai trò agent tự lợi:")
    bullet(doc, "GovernmentAgent — chính sách tài khóa, thuế, thuế quan.")
    bullet(doc, "CentralBankAgent — chính sách tiền tệ (ví dụ quy tắc kiểu Taylor), lãi suất, FX.")
    bullet(doc, "EnterpriseAgent — hành vi doanh nghiệp / chuỗi cung ứng.")
    bullet(doc, "PublicAgent — niềm tin tiêu dùng, bất ổn xã hội.")
    body(
        doc,
        "Các hub được nối bằng ma trận thương mại và lớp thuế quan. Đột biến cấu trúc (ví dụ tăng thuế quan "
        "giữa hai nước) lan truyền theo chuỗi nhân quả qua các tick mô phỏng; chronology fork lưu lịch sử "
        "kịch bản để tái quan sát.",
    )

    h2(doc, "6.3. Macro variables")
    body(
        doc,
        "Mỗi quốc gia mang vector biến số nhóm Monetary / Fiscal–Trade / Labor / Industrial / Geopolitical "
        "(lãi suất, CPI, tăng trưởng, thất nghiệp, thuế, dự trữ, chỉ số bất ổn…). Biến được clamp trong biên "
        "hợp lý để giữ ổn định số học của mô phỏng.",
    )

    h2(doc, "6.4. Phạm vi live hiện tại")
    body(
        doc,
        "Giao diện có thể hiển thị topology nhiều quốc gia; phần backend mutate live tập trung vào tập hub "
        "đã khai báo (ví dụ USA, CHN, VNM, JPN, IND, DEU và một số hub bổ sung). Đây là lựa chọn trung thực "
        "về phạm vi — ưu tiên độ sâu thí nghiệm hơn độ bao phủ toàn cầu.",
    )

    h2(doc, "6.5. Hypothesis Engine")
    bullet(doc, "Sinh giả thuyết shock (combinatorial mặc định; LLM tùy chọn, không bắt buộc cho thí nghiệm khoa học).")
    bullet(doc, "Debate / thinker hỗ trợ cấu trúc câu hỏi thí nghiệm.")
    bullet(doc, "Chạy scenario trên WorldKernel, đo pre/post macro deltas.")
    bullet(doc, "Ghi sealed JSONL rollout phục vụ phân tích và (sau này) dataset.")


def section_quant(doc):
    h1(doc, "07  ·  ECONITH QUANT")

    h2(doc, "7.1. Vai trò")
    quote(
        doc,
        "ECONITH Quant không phải trading bot. Đây là hệ thống AI nghiên cứu phản ứng của thị trường "
        "tài chính dưới các điều kiện khác nhau — bao gồm điều kiện stress được World tạo ra trong SIMULATION.",
    )

    h2(doc, "7.2. Pipeline")
    make_table(
        doc,
        ["Bước", "Mô tả", "Ghi chú kỹ thuật"],
        [
            ["Collect", "Thu thập market / macro / tradfi", "collectors/ — zero-ML, deploy độc lập"],
            ["Feature", "Đồng bộ đa tần số, chống look-ahead", "join_asof backward trong feature_pipeline"],
            ["Label", "Nhãn theo symbol, split thời gian", "label_symbol tránh cross-symbol contamination"],
            ["Train", "Huấn luyện desk (ví dụ PPO) khi có dữ liệu/GPU", "train_ppo; artifact không mặc định sẵn"],
            ["Infer", "Ensemble desk + regime → ai.signal", "Predictor; fallback heuristic nếu thiếu checkpoint"],
            ["Bridge", "Định cỡ / chuyển thành order.intent", "AIBridge + portfolio VaR haircut"],
            ["Risk", "Veto độc lập", "Sentinel (equity từ quant.fills)"],
            ["Evaluate", "Backtest offline, deploy gate", "evaluation/backtest, deploy_gate"],
        ],
    )

    h2(doc, "7.3. Quan hệ với mục tiêu nghiên cứu")
    body(
        doc,
        "Trong khung đề tài, Quant là công cụ đo lường và quan sát. Việc có desk heuristic hay đã train "
        "ảnh hưởng chất lượng tín hiệu, nhưng câu hỏi khoa học trung tâm vẫn là hiệu lực của bối cảnh "
        "macro mô phỏng — không phải tối ưu Sharpe trên vốn thật.",
    )


def section_coupling(doc):
    h1(doc, "08  ·  CƠ CHẾ KẾT NỐI WORLD ↔ QUANT")

    h2(doc, "8.1. Vì sao đây là phần quan trọng nhất?")
    body(
        doc,
        "Nếu chỉ có World hoặc chỉ có Quant, ECONITH gần với các hệ đã có. Điểm khác biệt phương pháp "
        "nằm ở lớp dịch shock hai chiều — cho phép thí nghiệm “macro context → market reaction” và "
        "quan sát phản hồi ngược trong SIMULATION. Đây cũng là phần dễ bị hiểu nhầm nhất: coupling "
        "không chứng minh dự báo chắc chắn, mà tạo điều kiện thí nghiệm có cấu trúc để so sánh với baseline.",
    )

    h2(doc, "8.2. Macro → Micro")
    body(doc, "Luồng khái niệm:", indent=False)
    bullet(doc, "Interest rate / tariff / inflation shock trên World")
    bullet(doc, "→ aggregate geo state")
    bullet(doc, "→ MicrostructuralVolatilityVector (vol multiplier, order-flow shock, liquidity drain, regime pressure)")
    bullet(doc, "→ Predictor điều chỉnh perception feature / regime (khi coupling được phép)")
    bullet(doc, "→ thay đổi tín hiệu AI Quant và metric quan sát")
    body(
        doc,
        "Ví dụ minh họa (stylized): thay đổi lãi suất → thanh khoản / kỳ vọng rủi ro thay đổi → "
        "biến động thị trường tăng → AI Quant điều chỉnh tín hiệu và Sentinel ghi nhận metric rủi ro. "
        "Chuỗi này được dùng để thiết kế thí nghiệm, không phải để khẳng định nhân quả tuyệt đối ngoài đời thực.",
    )

    h2(doc, "8.3. Micro → Macro")
    body(doc, "Luồng khái niệm:", indent=False)
    bullet(doc, "Market snapshot (conviction, vol, liquidation pressure…)")
    bullet(doc, "→ MacroFeedback theo mức mong manh cấu trúc từng quốc gia")
    bullet(doc, "→ cập nhật trạng thái World (ví dụ áp lực FX / unrest / yield)")
    bullet(doc, "→ trạng thái mô phỏng mới cho tick tiếp theo")

    h2(doc, "8.4. Ràng buộc an toàn nghiên cứu")
    bullet(doc, "Trong REALITY: EventBus có thể chặn world.* tới DOMAIN_QUANT; producer air-gap bổ sung.")
    bullet(doc, "Trong SIMULATION: cho phép World tác động Quant để phục vụ thí nghiệm.")
    bullet(doc, "Mục tiêu: tránh look-ahead / nhiễm mô phỏng vào kết luận “thực tế” khi không cố ý.")
    body(
        doc,
        "Bốn tầng cô lập (consumer gate, producer air-gap, execution air-gap, anomaly gate) là điều kiện "
        "kỹ thuật để báo cáo nghiên cứu phân biệt rõ “kết quả thí nghiệm mô phỏng” và “quan sát trên dữ liệu thật”.",
    )


def section_method(doc):
    h1(doc, "09  ·  PHƯƠNG PHÁP NGHIÊN CỨU & THỰC NGHIỆM")

    h2(doc, "9.1. Khung phương pháp")
    body(
        doc,
        "Áp dụng quy trình Research & Development (R&D) kết hợp thiết kế thí nghiệm so sánh có đối chứng. "
        "Biến độc lập tiêu biểu: có/không World coupling, loại shock (rate/tariff/inflation), mức severity, "
        "có/không calibration. Biến phụ thuộc tiêu biểu: sai số dự báo biến động, độ chính xác hướng, "
        "metric rủi ro, moment error, tỷ lệ tái lập.",
    )

    h2(doc, "9.2. Thiết kế thí nghiệm")
    make_table(
        doc,
        ["Nhánh", "Cấu hình", "Mục đích"],
        [
            ["B0 — Baseline", "Market data only; World coupling OFF", "Đối chứng tiêu chuẩn"],
            ["E1 — ECONITH", "Market data + simulated macro context (coupling ON)", "Điều kiện thí nghiệm"],
            ["C1 — Control", "Cùng magnitude nhưng shock ngẫu nhiên", "Loại trừ “đoán bừa”"],
        ],
    )
    body(
        doc,
        "Điều kiện kết luận tích cực tối thiểu: E1 cải thiện so với B0 trên metric đã công bố và đồng thời "
        "vượt C1 (random control). Nếu E1 ≤ C1, không được kết luận rằng mô phỏng “đúng”; chỉ kết luận "
        "rằng cấu trúc shock chưa mang lại tín hiệu vượt nhiễu trong thiết lập đó.",
    )

    h2(doc, "9.3. Quy trình thực nghiệm đề xuất")
    bullet(doc, "Bước 1 — Thu thập: market crypto + macro (FRED/…) + gắn nhãn sự kiện lịch sử.")
    bullet(doc, "Bước 2 — Calibration: moment-matching (OU + jump) → stochastic coefficients.")
    bullet(doc, "Bước 3 — Historical replay: inject shock khớp loại sự kiện; đo metric.")
    bullet(doc, "Bước 4 — Scenario testing: rate / inflation / tariff shocks với severity khác nhau.")
    bullet(doc, "Bước 5 — Ablation: World ON vs OFF; báo cáo bảng so sánh.")
    bullet(doc, "Bước 6 — Reproducibility: lặp fingerprint/seed; ghi tỷ lệ tái lập hướng delta.")

    h2(doc, "9.4. Metric")
    make_table(
        doc,
        ["Nhóm", "Metric ví dụ"],
        [
            ["World / Calibration", "Moment error (μ, σ, skew, kurtosis); direction accuracy trên case study"],
            ["Quant / Reaction", "MAE/RMSE biến động; directional accuracy của Δvol"],
            ["Risk", "Max drawdown; số lần Sentinel freeze / VaR breach (trong SIM)"],
            ["System", "Reproducibility rate; invariant isolation REALITY"],
        ],
    )

    h2(doc, "9.5. Lưu ý phương pháp")
    body(
        doc,
        "Đánh giá khoa học cần đo trên tín hiệu/policy thực của hệ thống. Các cổng kỹ thuật nội bộ "
        "(ví dụ gate deploy dùng tín hiệu giáo viên đơn giản) phục vụ vận hành kỹ thuật và không thay thế "
        "protocol thí nghiệm nêu trên. Thí nghiệm khoa học nên đặt HYPOTHESIS_USE_LLM=false để tăng tái lập.",
    )


def section_data(doc):
    h1(doc, "10  ·  DỮ LIỆU & THÀNH PHẦN KỸ THUẬT")

    h2(doc, "10.1. Nguồn dữ liệu")
    bullet(doc, "Market: order-flow / ticker crypto (collectors.market_coin).")
    bullet(doc, "Macro: chuỗi vĩ mô theo lịch (collectors.macro_global; FRED là nguồn chính khi có khóa).")
    bullet(doc, "TradFi tham chiếu: tài sản phiên giao dịch (collectors.tradfi_assets).")
    bullet(doc, "Event library (đề xuất nghiên cứu): COVID crash, chu kỳ tăng lãi suất, stress ngân hàng, tin thuế quan…")

    h2(doc, "10.2. Nguyên tắc dữ liệu")
    bullet(doc, "Raw lake append-only; feature store tách khỏi collectors.")
    bullet(doc, "Đồng bộ đa tần số không look-ahead (asof backward).")
    bullet(doc, "Label theo từng symbol để tránh nhiễm chéo timeline.")

    h2(doc, "10.3. Thành phần kỹ thuật chính")
    make_table(
        doc,
        ["Thành phần", "Chức năng ngắn"],
        [
            ["WorldKernel", "Máy trạng thái macro/micro theo tick"],
            ["SovereignWorldGraph", "Đồ thị quốc gia + 4 agent/hub"],
            ["Cross Impact", "Dịch shock hai chiều"],
            ["HypothesisRunner", "Vòng thí nghiệm giả thuyết"],
            ["Predictor", "Suy luận ensemble + regime"],
            ["AIBridge", "Tín hiệu → intent dưới ràng buộc rủi ro"],
            ["Sentinel", "Circuit breaker / VaR / equity truth"],
            ["Dashboard", "Quan sát /quant và /world"],
        ],
    )


def section_novelty(doc):
    h1(doc, "11  ·  TÍNH SÁNG TẠO")

    h2(doc, "11.1. Đóng góp thực chất")
    bullet(doc, "Kiến trúc tích hợp macro simulation và financial AI trong một event-driven runtime.")
    bullet(doc, "Mô hình tương tác hai chiều macro ↔ micro với vector shock có cấu trúc.")
    bullet(doc, "Hypothesis-driven simulation kèm sealed log phục vụ tái lập.")
    bullet(doc, "Môi trường thử nghiệm có ablation và isolation mode — hỗ trợ đánh giá AI dưới stress có chủ đích.")

    h2(doc, "11.2. Những gì không tính là sáng tạo khoa học")
    bullet(doc, "Việc “có AI”, “có dashboard”, hay “có nhiều module” đơn thuần.")
    bullet(doc, "Tích hợp thư viện bên thứ ba nếu không phục vụ câu hỏi nghiên cứu.")
    bullet(doc, "Claim marketing về tính “đầu tiên” khi thiếu bằng chứng so sánh có kiểm chứng.")


def section_applications(doc):
    h1(doc, "12  ·  ỨNG DỤNG THỰC TIỄN")
    bullet(doc, "Nghiên cứu tài chính–kinh tế: kiểm tra giả thuyết lan truyền shock.")
    bullet(doc, "Giáo dục: minh họa agent-based simulation và pipeline Quant trong một hệ thống.")
    bullet(doc, "Phân tích rủi ro / stress testing có chủ đích trên kịch bản.")
    bullet(doc, "Hỗ trợ ra quyết định nghiên cứu (decision support) — cung cấp bằng chứng thí nghiệm, không thay chuyên gia.")

    h2(doc, "12.1. Ranh giới sử dụng")
    body(
        doc,
        "ECONITH không được định vị là công cụ kiếm tiền tự động hay hệ thống thay thế phán đoán chuyên môn. "
        "Mọi suy luận từ mô phỏng cần đi kèm calibration, đối chứng baseline, và thừa nhận giới hạn stylized model. "
        "Trong bối cảnh báo cáo khoa học kỹ thuật, ứng dụng được hiểu là khả năng hỗ trợ quy trình nghiên cứu "
        "và học tập có kiểm chứng — không phải triển khai thương mại quy mô lớn.",
    )


def section_status(doc):
    h1(doc, "13  ·  TRẠNG THÁI HIỆN TẠI & GIỚI HẠN")

    h2(doc, "13.1. Đã có trong codebase")
    bullet(doc, "Runtime event-driven đầy đủ; dashboard Quant/World.")
    bullet(doc, "WorldKernel, Sovereign Graph, Cross Impact, Hypothesis Runner.")
    bullet(doc, "Collectors + feature/label pipeline + khung train/deploy.")
    bullet(doc, "Sentinel đồng bộ equity từ fills; cô lập REALITY/SIMULATION nhiều tầng.")
    bullet(doc, "Calibrator moment-matching (khung); paper soak checklist vận hành.")

    h2(doc, "13.2. Giới hạn cần nêu rõ")
    bullet(doc, "Desk AI mặc định có thể heuristic nếu chưa có checkpoint hợp lệ; không phải “AI đã train sẵn trong hộp”.")
    bullet(doc, "World mặc định có thể tắt trên máy yếu; coupling chỉ có hiệu lực nghiên cứu trong SIMULATION.")
    bullet(doc, "Calibration và historical validation là hướng phương pháp — cần chạy đủ thí nghiệm để thành bằng chứng.")
    bullet(doc, "Vòng tự động retrain→deploy (FULLY_AUTONOMOUS) là hướng vận hành tương lai, chưa phải năng lực hoàn tất.")
    bullet(doc, "Phạm vi hub live hữu hạn; không tương đương mô hình toàn cầu đầy đủ.")

    h2(doc, "13.3. Ý nghĩa của giới hạn đối với báo cáo")
    body(
        doc,
        "Việc nêu giới hạn không làm giảm giá trị đề tài; ngược lại, đây là điều kiện để hội đồng đánh giá đúng "
        "phần đã chứng minh bằng kiến trúc/code và phần còn cần số liệu thực nghiệm. Báo cáo này ưu tiên "
        "trung thực kỹ thuật hơn mở rộng tuyên bố.",
    )


def section_roadmap(doc):
    h1(doc, "14  ·  HƯỚNG PHÁT TRIỂN")

    h2(doc, "14.1. Ngắn hạn (ưu tiên khoa học)")
    bullet(doc, "Hoàn thiện calibration trên panel macro thật; báo cáo moment error.")
    bullet(doc, "Xây event library và chạy historical replay + random control.")
    bullet(doc, "Ablation có/không World với bảng metric công bố.")
    bullet(doc, "Mở rộng dataset market/macro dài ngày qua collectors.")

    h2(doc, "14.2. Trung hạn")
    bullet(doc, "Cải thiện fidelity agent và độ nhạy tham số (sensitivity analysis).")
    bullet(doc, "Đánh giá policy Quant thật trên holdout (không dùng tín hiệu oracle).")
    bullet(doc, "Chuẩn hóa báo cáo thí nghiệm / reproducibility kit.")

    h2(doc, "14.3. Dài hạn")
    bullet(doc, "Mở rộng quy mô mô phỏng và số thị trường được quan sát.")
    bullet(doc, "Nâng cấp agent kinh tế và cơ chế truyền dẫn phức tạp hơn — khi đã có validation nền.")
    bullet(doc, "Chỉ khi bằng chứng nghiên cứu vững mới cân nhắc vận hành DEMO kéo dài / tự động hóa vòng đời model.")


def section_conclusion(doc):
    h1(doc, "15  ·  KẾT LUẬN")
    body(
        doc,
        "ECONITH đóng góp một nền tảng kỹ thuật để nghiên cứu lan truyền shock vĩ mô tới phản ứng thị trường "
        "trong môi trường mô phỏng đa tác nhân có kết nối AI Quant. Giá trị cốt lõi nằm ở protocol thí nghiệm "
        "có kiểm soát — calibration, baseline, ablation, reproducibility — chứ không ở lời hứa dự đoán hay "
        "tự động hóa lợi nhuận.",
    )
    body(
        doc,
        "Với định vị này, dự án phù hợp làm hồ sơ nghiên cứu/kỹ thuật: trung thực về phạm vi, rõ về phương pháp, "
        "và mở đường cho các thực nghiệm kiểm chứng tiếp theo trên dữ liệu thật. Các bước tiếp theo ưu tiên "
        "bằng chứng số liệu (event library, bảng B0/E1/C1, moment error) hơn việc mở rộng tính năng giao diện.",
    )
    quote(
        doc,
        "ECONITH nghiên cứu vấn đề thiếu môi trường thí nghiệm đo được cho lan truyền shock vĩ mô → thị trường, "
        "bằng phương pháp agent-based simulation kết hợp AI Quant và so sánh baseline/ablation, "
        "nhằm hỗ trợ phân tích rủi ro và kiểm chứng giả thuyết trong điều kiện tái lập được.",
    )


def section_appendix(doc):
    h1(doc, "PHỤ LỤC A  ·  BẢN ĐỒ MODULE & THUẬT NGỮ")
    make_table(
        doc,
        ["Đường dẫn", "Vai trò"],
        [
            ["main.py", "Entrypoint FastAPI / lifespan runtime"],
            ["ai/simulator_engine/world_kernel.py", "World kernel"],
            ["ai/simulator_engine/sovereign_graph.py", "Sovereign multi-agent graph"],
            ["ai/simulator_engine/cross_impact.py", "Macro↔micro translators"],
            ["ai/simulator_engine/hypothesis_runner.py", "Hypothesis experiment loop"],
            ["ai/inference/predictor.py", "Quant inference ensemble"],
            ["econith_quant/bridge/ai_bridge.py", "Signal → order intent"],
            ["sentinel/manager.py", "Risk governance"],
            ["training/quant/feature_pipeline.py", "Multi-frequency features"],
            ["training/train_ppo.py", "Offline RL training entry"],
            ["collectors/", "Standalone data acquisition"],
            ["dashboard/", "Operator UI"],
            ["docs/paper_soak/", "Operational soak notes"],
        ],
    )

    h2(doc, "A.1. Thuật ngữ ngắn")
    bullet(doc, "SIMULATION — chế độ cho phép World coupling phục vụ thí nghiệm.")
    bullet(doc, "REALITY — chế độ ưu tiên dữ liệu thật; chặn nhiễm world vào order path.")
    bullet(doc, "Sealed rollout — bản ghi JSONL bất biến của một chu kỳ giả thuyết.")
    bullet(doc, "Ablation — tắt/bật thành phần để đo đóng góp trong thiết kế thí nghiệm.")
    bullet(doc, "Cross-impact — lớp dịch shock macro↔micro có cấu trúc.")
    bullet(doc, "Sentinel — lớp rủi ro độc lập, có quyền phủ quyết tín hiệu.")


def build():
    doc = Document()
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.0
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

    setup_page(doc)

    # Page 1 — cover
    cover_page(doc)
    page_break(doc)

    # Remaining content flows continuously; Word will paginate to ~15
    toc_page(doc)
    section_summary(doc)
    section_overview(doc)
    section_problem(doc)
    section_goals(doc)
    section_architecture(doc)
    section_world(doc)
    section_quant(doc)
    section_coupling(doc)
    section_method(doc)
    section_data(doc)
    section_novelty(doc)
    section_applications(doc)
    section_status(doc)
    section_roadmap(doc)
    section_conclusion(doc)
    section_appendix(doc)

    doc.save(OUT)
    print(f"Wrote {OUT}")
    return OUT


def count_pages(path: Path) -> int:
    import win32com.client

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(str(path.resolve()), ReadOnly=True)
        doc.Repaginate()
        pages = int(doc.ComputeStatistics(2))  # wdStatisticPages = 2
        doc.Close(False)
        return pages
    finally:
        word.Quit()


if __name__ == "__main__":
    out = build()
    try:
        n = count_pages(out)
        print(f"PAGE_COUNT={n}")
    except Exception as exc:  # noqa: BLE001
        print(f"PAGE_COUNT_ERR={type(exc).__name__}: {exc}")
